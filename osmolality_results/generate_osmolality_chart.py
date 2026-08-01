#!/usr/bin/env python3
"""Generate an Excel-style osmolality results chart with specification ranges."""

from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.lines import Line2D
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).resolve().parent

Y_MIN = 200
Y_MAX = 620
Y_STEP = 25

PRODUCTS = [
    {
        "product": "DO1",
        "laboratorio": "Laboratorio B",
        "spec_min": 280,
        "spec_max": 320,
        "sample_1": 282,
        "sample_2": 280,
        "asterisk": None,
    },
    {
        "product": "DO2",
        "laboratorio": "Laboratorio E",
        "spec_min": 240,
        "spec_max": 340,
        "sample_1": 286,
        "sample_2": 283,
        "asterisk": None,
    },
    {
        "product": "DO3",
        "laboratorio": "Laboratorio F",
        "spec_min": 250,
        "spec_max": 350,
        "sample_1": 280,
        "sample_2": 281,
        "asterisk": None,
    },
    {
        "product": "DO4",
        "laboratorio": "Laboratorio G",
        "spec_min": 260,
        "spec_max": 320,
        "sample_1": 281,
        "sample_2": 276,
        "asterisk": None,
    },
    {
        "product": "DO5",
        "laboratorio": "Laboratorio G",
        "spec_min": 260,
        "spec_max": 320,
        "sample_1": 276,
        "sample_2": 275,
        "asterisk": None,
    },
    {
        "product": "DO6",
        "laboratorio": "Laboratorio H",
        "spec_min": 240,
        "spec_max": 370,
        "sample_1": 287,
        "sample_2": 280,
        "asterisk": None,
    },
    {
        "product": "DO7",
        "laboratorio": "Laboratorio I",
        "spec_min": 270,
        "spec_max": 330,
        "sample_1": 310,
        "sample_2": 312,
        "asterisk": None,
    },
    {
        "product": "DO8",
        "laboratorio": "Laboratorio J",
        "spec_min": None,
        "spec_max": None,
        "sample_1": 573,
        "sample_2": 579,
        "asterisk": 574,
    },
    {
        "product": "DO9",
        "laboratorio": "Laboratorio K",
        "spec_min": None,
        "spec_max": None,
        "sample_1": None,
        "sample_2": None,
        "asterisk": 288,
    },
]


def build_chart_png(path: Path) -> None:
    """Excel-like chart: shaded spec ranges + circular lote markers + asterisks."""
    products = [p["product"] for p in PRODUCTS]
    x = np.arange(len(products))

    fig, ax = plt.subplots(figsize=(13, 8.0), dpi=160)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")

    bar_width = 0.55
    for i, p in enumerate(PRODUCTS):
        if p["spec_min"] is None or p["spec_max"] is None:
            continue
        height = p["spec_max"] - p["spec_min"]
        ax.bar(
            i,
            height,
            bottom=p["spec_min"],
            width=bar_width,
            color="#9DC3E6",
            alpha=0.45,
            edgecolor="#5B9BD5",
            linewidth=0.8,
            zorder=1,
        )

    offset = 0.08
    x1, y1 = [], []
    x2, y2 = [], []
    xa, ya = [], []
    for i, p in enumerate(PRODUCTS):
        if p["sample_1"] is not None:
            x1.append(i - offset)
            y1.append(p["sample_1"])
        if p["sample_2"] is not None:
            x2.append(i + offset)
            y2.append(p["sample_2"])
        if p["asterisk"] is not None:
            xa.append(i)
            ya.append(p["asterisk"])

    ax.scatter(
        x1,
        y1,
        s=55,
        facecolors="#ED7D31",
        edgecolors="#C55A11",
        linewidths=1.0,
        zorder=3,
        label="Lote 1",
    )
    ax.scatter(
        x2,
        y2,
        s=55,
        facecolors="#70AD47",
        edgecolors="#548235",
        linewidths=1.0,
        zorder=3,
        label="Lote 2",
    )
    ax.scatter(
        xa,
        ya,
        s=120,
        marker="*",
        facecolors="#000000",
        edgecolors="#000000",
        linewidths=0.4,
        zorder=4,
        label="Valor teórico estimado",
    )
    ax.axhline(
        y=300,
        color="#000080",
        linestyle="--",
        linewidth=1.4,
        zorder=2,
        label="Lágrima natural",
    )

    ax.set_axisbelow(True)
    ax.yaxis.grid(True, which="major", color="#D9D9D9", linestyle="-", linewidth=0.8)
    ax.xaxis.grid(True, which="major", color="#F2F2F2", linestyle="-", linewidth=0.6)
    ax.set_xlim(-0.6, len(products) - 0.4)
    ax.set_ylim(Y_MIN, Y_MAX)
    ax.set_yticks(np.arange(Y_MIN, Y_MAX + 1, Y_STEP))
    ax.set_xticks(x)
    ax.set_xticklabels(products, fontsize=11)
    ax.set_ylabel("Osmolalidad (mOsm/kg)", fontsize=11)
    ax.set_title(
        "Osmolality Test Results by Product",
        fontsize=14,
        fontweight="bold",
        pad=14,
        color="#1F4E79",
    )

    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    ax.spines["left"].set_color("#BFBFBF")
    ax.spines["bottom"].set_color("#BFBFBF")
    ax.tick_params(colors="#595959")
    ax.tick_params(axis="x", pad=4)

    # Laboratory names horizontal + bold under product ticks (grouped when shared)
    lab_groups = []
    for i, p in enumerate(PRODUCTS):
        lab = p["laboratorio"]
        if lab_groups and lab_groups[-1]["name"] == lab:
            lab_groups[-1]["end"] = i
        else:
            lab_groups.append({"name": lab, "start": i, "end": i})

    for group in lab_groups:
        center = (group["start"] + group["end"]) / 2
        ax.text(
            center,
            -0.078,
            group["name"],
            transform=ax.get_xaxis_transform(),
            ha="center",
            va="top",
            fontsize=8,
            fontweight="bold",
            color="#404040",
            clip_on=False,
        )
        if group["start"] != group["end"]:
            ax.annotate(
                "",
                xy=(group["end"] + 0.35, -0.052),
                xytext=(group["start"] - 0.35, -0.052),
                xycoords=("data", "axes fraction"),
                textcoords=("data", "axes fraction"),
                arrowprops=dict(
                    arrowstyle="-",
                    color="#BFBFBF",
                    lw=0.9,
                    shrinkA=0,
                    shrinkB=0,
                ),
                annotation_clip=False,
            )

    ax.set_xlabel("Laboratorio/Producto", fontsize=11, labelpad=36)

    legend_handles = [
        mpatches.Patch(
            facecolor="#9DC3E6",
            edgecolor="#5B9BD5",
            alpha=0.45,
            label="Rango de aceptación",
        ),
        Line2D(
            [0],
            [0],
            marker="o",
            color="w",
            markerfacecolor="#ED7D31",
            markeredgecolor="#C55A11",
            markersize=8,
            label="Lote 1",
        ),
        Line2D(
            [0],
            [0],
            marker="o",
            color="w",
            markerfacecolor="#70AD47",
            markeredgecolor="#548235",
            markersize=8,
            label="Lote 2",
        ),
        Line2D(
            [0],
            [0],
            marker="*",
            color="w",
            markerfacecolor="#000000",
            markeredgecolor="#000000",
            markersize=12,
            label="Valor teórico estimado",
        ),
        Line2D(
            [0],
            [0],
            color="#000080",
            linestyle="--",
            linewidth=1.6,
            label="Lágrima natural",
        ),
    ]
    ax.legend(handles=legend_handles, loc="upper left", frameon=True, fancybox=False)

    fig.subplots_adjust(left=0.10, right=0.98, top=0.92, bottom=0.20)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def within_spec(value, lo, hi) -> str:
    if value is None:
        return "N/A (sin lote)"
    if lo is None or hi is None:
        return "N/A (no specification)"
    return "Pass" if lo <= value <= hi else "Out of range"


def build_workbook(chart_png: Path, path: Path) -> None:
    wb = Workbook()
    n = len(PRODUCTS)
    last_data_row = 1 + n

    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum["A1"] = "Osmolality Test — Summary"
    ws_sum["A1"].font = Font(bold=True, size=16, color="1F4E79")
    ws_sum["A3"] = (
        "All products with a specification range are within limits for both lots."
    )
    ws_sum["A4"] = (
        "DO8 has no specification (lotes ~573–579; asterisk at 574). "
        "DO9 (Laboratorio K) has asterisk-only value at 288 mOsm/kg."
    )
    ws_sum["A6"] = "Contents"
    ws_sum["A6"].font = Font(bold=True)
    ws_sum["A7"] = "• Sheet 'Data': raw results + native Excel combo chart"
    ws_sum["A8"] = "• Sheet 'Chart': high-resolution Excel-style chart image"
    ws_sum["A9"] = "• Companion PNG file for presentations / reports"
    ws_sum.column_dimensions["A"].width = 95

    ws = wb.create_sheet("Data")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="B4B4B4"),
        right=Side(style="thin", color="B4B4B4"),
        top=Side(style="thin", color="B4B4B4"),
        bottom=Side(style="thin", color="B4B4B4"),
    )
    center = Alignment(horizontal="center", vertical="center")

    headers = [
        "Product",
        "Laboratorio",
        "Spec Min (mOsm/kg)",
        "Spec Max (mOsm/kg)",
        "Lote 1 (mOsm/kg)",
        "Lote 2 (mOsm/kg)",
        "Asterisk (mOsm/kg)",
        "Lote 1 Status",
        "Lote 2 Status",
        "Base (hidden)",
        "Range Height",
        "X Index",
        "Lote 1 Y",
        "Lote 2 Y",
        "Asterisk Y",
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    for row_idx, p in enumerate(PRODUCTS, start=2):
        values = [
            p["product"],
            p["laboratorio"],
            p["spec_min"] if p["spec_min"] is not None else "—",
            p["spec_max"] if p["spec_max"] is not None else "—",
            p["sample_1"] if p["sample_1"] is not None else "—",
            p["sample_2"] if p["sample_2"] is not None else "—",
            p["asterisk"] if p["asterisk"] is not None else "—",
            within_spec(p["sample_1"], p["spec_min"], p["spec_max"]),
            within_spec(p["sample_2"], p["spec_min"], p["spec_max"]),
            p["spec_min"] if p["spec_min"] is not None else 0,
            (p["spec_max"] - p["spec_min"]) if p["spec_min"] is not None else 0,
            row_idx - 1,
            p["sample_1"],
            p["sample_2"],
            p["asterisk"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col, value if value is not None else None)
            cell.alignment = center
            cell.border = thin
            if col in (8, 9):
                if value == "Pass":
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
                    cell.font = Font(color="006100")
                elif str(value).startswith("Out"):
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
                    cell.font = Font(color="9C0006")
                else:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                    cell.font = Font(color="9C5700")

    widths = [12, 16, 18, 18, 18, 18, 18, 22, 22, 14, 14, 10, 12, 12, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for col in range(10, 16):
        ws.column_dimensions[get_column_letter(col)].hidden = True

    notes_row = last_data_row + 2
    ws.cell(notes_row, 1, "Notes").font = Font(bold=True, color="1F4E79")
    ws.cell(
        notes_row + 1,
        1,
        "Rango de aceptación shown as a shaded band. "
        "Asterisk (*) markers: DO8 = 574, DO9 = 288. "
        "X-axis shows product and belonging Laboratorio.",
    )
    ws.cell(
        notes_row + 2,
        1,
        "Units: mOsm/kg. Markers: Lote 1 (orange), Lote 2 (green), "
        "estrella = Valor teórico estimado; línea punteada azul marino = Lágrima natural (300).",
    )
    ws.merge_cells(
        start_row=notes_row + 1, start_column=1, end_row=notes_row + 1, end_column=9
    )
    ws.merge_cells(
        start_row=notes_row + 2, start_column=1, end_row=notes_row + 2, end_column=9
    )

    bar = BarChart()
    bar.type = "col"
    bar.grouping = "stacked"
    bar.overlap = 100
    bar.title = "Osmolality Test Results by Product"
    bar.y_axis.title = "Osmolalidad (mOsm/kg)"
    bar.x_axis.title = "Laboratorio/Producto"
    bar.y_axis.scaling.min = Y_MIN
    bar.y_axis.scaling.max = Y_MAX
    bar.y_axis.majorUnit = Y_STEP
    bar.style = 10
    bar.width = 20
    bar.height = 12

    cats = Reference(ws, min_col=1, min_row=2, max_row=last_data_row)
    data_bar = Reference(ws, min_col=10, min_row=1, max_col=11, max_row=last_data_row)
    bar.add_data(data_bar, titles_from_data=True)
    bar.set_categories(cats)

    bar.series[0].graphicalProperties.noFill = True
    bar.series[0].graphicalProperties.line.noFill = True
    bar.series[0].title = SeriesLabel(v=" ")

    bar.series[1].graphicalProperties.solidFill = "9DC3E6"
    bar.series[1].graphicalProperties.line.solidFill = "5B9BD5"
    bar.series[1].title = SeriesLabel(v="Rango de aceptación")

    scatter = ScatterChart()
    scatter.style = 10

    xvalues = Reference(ws, min_col=12, min_row=2, max_row=last_data_row)
    y1_vals = Reference(ws, min_col=13, min_row=2, max_row=last_data_row)
    y2_vals = Reference(ws, min_col=14, min_row=2, max_row=last_data_row)
    ya_vals = Reference(ws, min_col=15, min_row=2, max_row=last_data_row)

    ser1 = Series(y1_vals, xvalues, title="Lote 1")
    ser1.marker = Marker(symbol="circle", size=7)
    ser1.marker.graphicalProperties.solidFill = "ED7D31"
    ser1.marker.graphicalProperties.line.solidFill = "C55A11"
    ser1.graphicalProperties.line.noFill = True
    scatter.series.append(ser1)

    ser2 = Series(y2_vals, xvalues, title="Lote 2")
    ser2.marker = Marker(symbol="circle", size=7)
    ser2.marker.graphicalProperties.solidFill = "70AD47"
    ser2.marker.graphicalProperties.line.solidFill = "548235"
    ser2.graphicalProperties.line.noFill = True
    scatter.series.append(ser2)

    ser_a = Series(ya_vals, xvalues, title="Valor teórico estimado")
    ser_a.marker = Marker(symbol="star", size=10)
    ser_a.marker.graphicalProperties.solidFill = "000000"
    ser_a.marker.graphicalProperties.line.solidFill = "000000"
    ser_a.graphicalProperties.line.noFill = True
    scatter.series.append(ser_a)

    bar.y_axis.majorGridlines.spPr = None
    bar += scatter
    ws.add_chart(bar, f"A{notes_row + 4}")

    ws_chart = wb.create_sheet("Chart")
    ws_chart["A1"] = "Osmolality Test Results — Excel-style chart"
    ws_chart["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws_chart["A2"] = (
        "Eje X: Producto (DO1–DO9) + Laboratorio | Eje Y: Osmolalidad (mOsm/kg), paso 25 | "
        "Sombra: Rango de aceptación | Círculos: Lote 1 & Lote 2 | "
        "Estrella: Valor teórico estimado | Línea punteada: Lágrima natural (300)"
    )
    ws_chart.merge_cells("A2:H2")
    ws_chart.column_dimensions["A"].width = 20

    if chart_png.exists():
        img = XLImage(str(chart_png))
        img.width = 1020
        img.height = 580
        ws_chart.add_image(img, "A4")

    wb.save(path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    png_path = OUTPUT_DIR / "osmolality_chart.png"
    xlsx_path = OUTPUT_DIR / "osmolality_results.xlsx"

    build_chart_png(png_path)
    build_workbook(png_path, xlsx_path)
    print(f"Wrote {png_path}")
    print(f"Wrote {xlsx_path}")


if __name__ == "__main__":
    main()
